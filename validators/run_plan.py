"""
Everything the workbook still needs, at once — and the contract an assistant reads.

THE DEFECT THIS ANSWERS
------------------------
`load_config` raises on the FIRST problem. That is right for a gate and wrong
for a person. On 2026-09-06 the owner ran a fresh template and was told the
table file did not exist; fixing that would have got him "no split", then "no
key", then "the key is not defined" — four sittings to learn what one screen
can say. He is an economist who does not use a terminal, which is exactly the
user `docs/GUIDE.md` opens by promising to serve.

`--plan` reads the workbook leniently and reports every gap in one pass, each
with why it matters and what to put. On a workbook broken four ways it prints
six blocking problems and two weaknesses where `--check` prints one error.

AND WHY IT IS NOT A CHATBOT
-----------------------------
The owner asked on 2026-09-07 for natural language to reach the workbook: say
what you want, get a configuration back to approve. A language model cannot
live inside this engine — `pyproject.toml` promises numpy and openpyxl, the run
must be deterministic, and a workbook has to give the same numbers on any day
for the DOI to mean anything.

So the understanding happens OUTSIDE and the engine owes the other half of the
contract: an exact, machine-readable statement of what it needs. `--plan
--json` is that statement. An assistant reads it, fills the sheet, and the
USER approves it before anything runs. Nothing in `--plan` computes, fetches or
writes, and that is checked here rather than asserted.

WHAT IT MUST NOT DO
---------------------
Contradict itself. The first version reported "no allocation key is registered"
and "only one allocation key is registered" in the same screen, which is how a
reader stops believing a report.

Exit non-zero. `--plan` is advice, and advice that exits 1 is the defect
`--sources` carried until v1.85, printing "Nothing is wrong" while telling the
shell otherwise. `--check` is the gate; the REPORT says whether the workbook is
ready.
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

FAIL = []


def check(label, ok, detail=""):
    print(f"  [{'ok' if ok else 'FAIL'}] {label}")
    if detail:
        print(f"         {detail}")
    if not ok:
        FAIL.append(label)


def main():
    import openpyxl

    from quadrium.config import plan_workbook, write_template

    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        cfg = write_template(td / "t.xlsx")

        # 1 -- an untouched template: one blocker, and it is its own seed.
        rep = plan_workbook(cfg)
        blocking = [g for g in rep["gaps"] if g["severity"] == "blocking"]
        check("an untouched template is not ready", rep["ready"] is False)
        check("and the one thing stopping it is the value it seeded itself",
              len(blocking) == 1
              and "--template wrote" in blocking[0]["what"],
              f"{len(blocking)} blocker(s): "
              f"{[g['what'] for g in blocking]}")
        check("and the template's invented keys are flagged as weak, not "
              "as errors",
              any("REPLACE" in g["what"] for g in rep["gaps"]
                  if g["severity"] == "weak"),
              "they run; what they produce is a demonstration")

        # 2 -- broken four ways, reported at once.
        wb = openpyxl.load_workbook(cfg)
        pj = wb["project"]
        pj["B2"] = "/no/such/file.xlsx"
        pj["B3"] = "ine_interor"                       # a typo, not a kind
        pj["A6"], pj["B6"] = "table_unbalanced", "residual_column"

        def wipe(name, n):
            ws = wb[name]
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if isinstance(v, str) and v.lstrip().startswith("#"):
                    continue
                for c in range(1, n + 1):
                    ws.cell(r, c).value = None
            return ws

        sp = wipe("splits", 4)
        sp["A2"], sp["B2"], sp["D2"] = "36", "36A", "k36"
        wipe("keys", 6)
        wipe("profiles", 4)
        wb.save(cfg)

        rep = plan_workbook(cfg)
        blocking = [g for g in rep["gaps"] if g["severity"] == "blocking"]
        found = " | ".join(g["what"] for g in blocking)
        check("four independent faults are all reported in one pass",
              len(blocking) >= 5,
              f"{len(blocking)} blockers, where --check would raise on one")
        for want, why in (
                ("is not a kind", "the misspelt table_kind"),
                ("is not there", "the missing file"),
                ("1 subsector", "a split into one piece"),
                ("no allocation key", "the empty keys sheet"),
                ("does not define it", "the key named and never defined")):
            check(f"it names {why}", want in found)

        # 3 -- and it does not contradict itself.
        check("with no keys at all it does not also say there is one",
              not any("only one allocation key" in g["what"]
                      for g in rep["gaps"]),
              "reported beside 'no allocation key is registered' this "
              "contradicted it in the same screen")

        # 4 -- a file that is not a workbook is described, not crashed on.
        bad = td / "notaworkbook.xlsx"
        bad.write_text("this is not a spreadsheet")
        rep = plan_workbook(bad)
        check("a file that cannot be opened is reported, not raised",
              rep["ready"] is False and rep["gaps"]
              and "cannot be opened" in rep["gaps"][0]["what"],
              "the one thing --plan cannot survive, and it says so")

        # 5 -- the machine-readable half, through the command line, and the
        #      exit code that says advice rather than verdict.
        run = subprocess.run(
            [sys.executable, str(ROOT / "run_quadrium.py"), str(cfg),
             "--plan", "--json"],
            capture_output=True, text=True, cwd=td)
        check("`--plan --json` exits 0 even when the workbook is not ready",
              run.returncode == 0,
              f"exit {run.returncode} — advice that exits non-zero is the "
              f"defect --sources carried until v1.85")
        try:
            doc = json.loads(run.stdout)
        except ValueError:
            doc = None
        check("and prints JSON an assistant can act on",
              isinstance(doc, dict)
              and {"ready", "gaps", "have"} <= set(doc)
              and all({"sheet", "what", "why", "fix", "severity"} <= set(g)
                      for g in doc["gaps"]),
              "every gap carries where it is, what it is, why it matters and "
              "what to put — which is what a model needs to fill the sheet")

        # 6 -- the PROSE path, which is the one a person actually meets.
        #      Checking only --json would leave the human printer entered by
        #      nothing, which is how this project keeps finding features that
        #      were built, verified and unreachable.
        human = subprocess.run(
            [sys.executable, str(ROOT / "run_quadrium.py"), str(cfg),
             "--plan"], capture_output=True, text=True, cwd=td)
        check("`--plan` in prose runs and reads as a list, not as an error",
              human.returncode == 0
              and "stop it running" in human.stdout
              and "why:" in human.stdout and "put:" in human.stdout,
              f"exit {human.returncode}")
        check("and it says plainly that it touched nothing",
              "Nothing was computed, fetched or written" in human.stdout)

        # 7 -- it reads and does nothing else.
        before = {p.name: p.stat().st_mtime_ns for p in td.iterdir()}
        plan_workbook(cfg)
        after = {p.name: p.stat().st_mtime_ns for p in td.iterdir()}
        check("planning writes nothing and changes nothing",
              before == after,
              "it describes a file; a planner that edits is not a planner")

    print()
    print("    --check is the gate and stops at the first problem. This is")
    print("    the list, and the list is what a person -- or an assistant --")
    print("    needs to finish the job in one sitting.")

    print("\n" + "=" * 78)
    if FAIL:
        print(f"{len(FAIL)} check(s) FAILED: {', '.join(FAIL)}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
