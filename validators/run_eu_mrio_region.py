"""
The European MRIO as one region's table: the archive does not balance, and the
residue is carried where it can be seen.

WHAT WAS OPEN
---------------
`INDEX.md` §7 A1 left one thing unclosed. The 272-region MRIO of Huang and
Koutroumpis (2023) does not balance, and building an `IOTable` from it meant
choosing between deriving output from the identity, carrying the residue in a
named column, and refusing the 123 units whose implied final demand is
negative. The owner chose on 2026-09-11, on three measurements that
`run_mrio_real_output.py` holds:

    derive output        multipliers move 8.1 % at the median, 198 % at worst,
                         and negative value added goes from 26 units to 96
    refuse the 123       they sit in 83 of the 272 regions
    keep published X     multipliers are exactly the archive's

**The third, because a multiplier reads only `Z` and `X`.** Both are published
and they are one vector: the final-demand file's `TOTAL` and the value-added
file's `INPUT` agree unit by unit to the last digit, so the objection that
refused Austria's tables -- two output vectors and an `IOTable` needs one --
does not arise here. What does not close is the COMPONENTS: the residue lives
in `Y` and `VA`, where levels are read, and never in `A`.

AND THE 123 WERE NEVER THE PROBLEM
------------------------------------
They were the part of it that showed. The row residue is non-zero in every
unit that has any output -- 2,680 of 2,720, the other 40 being the four empty
regions -- and at the median it is 16 % of a region's output. The 123 are only
the units where it is large enough to flip the sign of implied final demand.
Refusing them would have treated the symptom and kept the cause everywhere
else. (The plan that preceded this file said "non-zero in 2,184"; that is the
number of NEGATIVE residues, and `run_mrio_real_output.py` now measures both.)

WHAT THIS CHECKS
------------------
One region, loaded through the engine, against the same numbers read
independently through `run_mrio_axis_scale.py`'s own readers: the block is the
region's diagonal block, output is the published `TOTAL`, both identities close
because the residue is carried in ONE column and ONE row that say what they
are, the residue is the archive's and not an invention, the coefficients are
the archive's, and the share of the multiplier this single-region table omits
matches the full 2,720 x 2,720 inverse. Then every refusal is fired, and the
whole thing is run from a workbook by someone who does not program.

Run:
    python3 validators/run_eu_mrio_region.py
"""
from __future__ import annotations

import importlib.util
import io
import sys
import tempfile
import warnings
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
MRIO = ROOT / "data" / "mrio"
FILES = ("MRIO_2018_272regions.xlsx", "Final_demand_2018.xlsx",
         "TAXSUB_VA_2018.xlsx")
REGION = "ES51"
S = 10
FAIL: list[str] = []

# Exit code 3, not 0. `check.sh` counts it as SKIPPED rather than as a pass:
# this validator opened no file and measured nothing, and a run that says
# "All checks passed" on evidence it never saw is the fault the suite exists
# to catch. It is not a failure -- a tree without the workbook still exits 0.
NOTHING_CHECKED = 3


def check(name: str, ok: bool, detail: str = "") -> None:
    print(f"  {'ok  ' if ok else 'FAIL'} {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        FAIL.append(name)


def run(argv):
    from quadrium.cli import main

    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        code = main(argv)
    return code, out.getvalue(), err.getvalue()


def workbook(path, meta, *, splits=True, regionalise=None):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "project"
    for k, v in meta:
        ws.append([k, v])
    if regionalise is not None:
        ws = wb.create_sheet("regionalise")
        ws.append(["key", "value"])
        for row in regionalise:
            ws.append(list(row))
    if splits:
        ws = wb.create_sheet("splits")
        ws.append(["sector_code", "new_code", "new_label", "key_id"])
        ws.append(["G-I", "GI1", "Trade and transport", "k1"])
        ws.append(["G-I", "GI2", "Accommodation and food", "k1"])
        ws = wb.create_sheet("keys")
        ws.append(["key_id", "new_sector_code", "value", "source",
                   "source_year", "strength"])
        ws.append(["k1", "GI1", 70, "validator fixture, not a measurement",
                   2018, "weak"])
        ws.append(["k1", "GI2", 30, "validator fixture, not a measurement",
                   2018, "weak"])
    wb.save(path)
    return path


def refused(fn, word):
    """(True, first line) when `fn` raises and the message names `word`."""
    try:
        fn()
    except Exception as exc:                              # noqa: BLE001
        msg = str(exc)
        return word in msg, msg.strip().splitlines()[0][:150]
    return False, "it loaded"


def main() -> int:
    warnings.filterwarnings("ignore")
    print(__doc__.strip().split("Run:")[0].rstrip())
    print("\n" + "=" * 78)

    missing = [f for f in FILES if not (MRIO / f).exists()]
    if missing:
        print(f"    -- {', '.join(missing)} absent (gitignored). The URL and")
        print("       SHA-256 of the archive are in data/mrio/_provenance.json.")
        print("\n" + "=" * 78)
        print("Nothing was checked: the MRIO archive is not in this tree.")
        return NOTHING_CHECKED

    try:
        from quadrium.io_loader import load_eu_mrio_2018
    except ImportError as exc:
        check("the engine has a loader for the archive", False, str(exc))
        print(f"\n{len(FAIL)} check(s) FAILED.")
        return 1

    # ---- the same numbers, read independently of the loader
    spec = importlib.util.spec_from_file_location(
        "axis", ROOT / "validators" / "run_mrio_axis_scale.py")
    axis = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(axis)
    Z, labels = axis.load_Z()
    _, fh, FD = axis.load_side(axis.FD, "rows")
    _, vh, VA = axis.load_side(axis.VA, "columns")
    regions = list(dict.fromkeys(l.split("-", 1)[0] for l in labels))
    k = regions.index(REGION)
    s = slice(k * S, (k + 1) * S)
    X = FD[:, fh.index("TOTAL")]

    t = load_eu_mrio_2018(MRIO, REGION)

    check("one region's own table, with the archive's published output",
          t.n == S and np.array_equal(t.X, X[s]),
          f"{REGION}: {t.n} sectors, output {t.X.sum():,.0f}, identical to the "
          f"final-demand file's TOTAL. Nothing is derived")
    check("its intermediate block is the region's diagonal block, cell for "
          "cell", np.array_equal(t.Z, Z[s, s]),
          "the block and the side files join by POSITION, the correspondence "
          "run_mrio_side_join.py established; the labels are not read")

    row = t.Z.sum(1) + t.Y.sum(1) - t.X
    col = t.Z.sum(0) + t.VA.sum(0) - t.X
    floor = 1e-9 * float(t.X.max())
    check("both identities close", float(np.abs(row).max()) < floor
          and float(np.abs(col).max()) < floor,
          f"worst row {np.abs(row).max():.2e}, worst column "
          f"{np.abs(col).max():.2e} — because the residue is carried, not "
          f"because anything was adjusted")

    iy = [i for i, l in enumerate(t.Y_labels) if "RESIDUAL" in l]
    iv = [i for i, l in enumerate(t.VA_labels) if "RESIDUAL" in l]
    check("in one column and one row that say what they are",
          len(iy) == 1 and len(iv) == 1,
          f"{t.Y_labels[iy[0]] if iy else '(no residual column)'!r}")

    keep = [fh.index(h) for h in ("HFCE", "GGFC", "GFCF", "INVNT", "EX")]
    ry = X[s] - Z[s].sum(1) - FD[s][:, keep].sum(1)
    v3 = VA[s][:, [vh.index(h) for h in ("TAXSUB", "VA", "IM")]].sum(1)
    rv = X[s] - Z[:, s].sum(0) - v3
    same = bool(iy and iv and np.allclose(t.Y[:, iy[0]], ry, rtol=0, atol=1e-6)
                and np.allclose(t.VA[iv[0]], rv, rtol=0, atol=1e-6))
    share_y = float(np.abs(ry).sum() / X[s].sum())
    share_v = float(np.abs(rv).sum() / X[s].sum())
    check("and the residue is the archive's own, not something the loader "
          "chose", same,
          f"recomputed from the three files without the loader: row "
          f"{100 * share_y:.1f} % and column {100 * share_v:.1f} % of the "
          f"region's output, in absolute value")
    head = (t.notes or "")[:400]
    check("the notes say how large it is before anything else",
          f"{100 * share_y:.1f} %" in head and f"{100 * share_v:.1f} %" in head,
          "reporting.py prints the notes under 'What the loader decided', "
          "above every number in report.md")

    A = Z / np.where(X > 0, X, np.inf)
    A[~np.isfinite(A)] = 0.0
    check("the coefficients are the archive's, so the residue reaches no "
          "multiplier", np.allclose(t.Z / t.X, A[s, s], rtol=1e-12, atol=0),
          "A = Z / X reads the published block and the published output and "
          "nothing else. That is the reason the residue could be carried "
          "rather than removed")

    L = np.linalg.inv(np.eye(len(A)) - A)
    m = L[:, s].sum(0)
    intra = L[s, s].sum(0)
    agg = float((m - intra).sum() / m.sum())
    check("the notes carry what a single-region table omits, from the full "
          "inverse", f"{100 * agg:.1f} %" in (t.notes or ""),
          f"{100 * agg:.1f} % of {REGION}'s output multipliers travels through "
          f"other regions and comes back. No regionalisation by quotients can "
          f"give that number; it is the reason this source is worth loading")

    ir = getattr(t, "interregional", None) or {}
    leak = {r: float(L[j * S:(j + 1) * S, s].sum())
            for j, r in enumerate(regions) if r != REGION}
    top = max(leak, key=leak.get)
    per_ind = (m - intra) / m
    check("and the table carries it sector by sector, not only as a sentence",
          bool(ir)
          and np.allclose(ir.get("share_by_sector", []), per_ind, atol=1e-9)
          and abs(ir.get("share", -1.0) - agg) < 1e-12
          and (ir.get("to_regions") or [[None]])[0][0] == top,
          f"{len(ir.get('share_by_sector', []))} sector shares, each equal to "
          f"the independent inverse; most of what leaves {REGION} goes to "
          f"{top}, {100 * leak[top] / float((m - intra).sum()):.1f} % of it")

    # ---- and what the region would lose at the surveys' level of trade
    #
    # The same counterfactual as the archive-wide range in
    # run_spillover_sensitivity.py -- trade with the rest of the country
    # multiplied by the factor, every column total held -- computed here with
    # THAT file's own implementation, so the loader's figure is checked
    # against code it does not share.
    from quadrium.regionalise import EVIDENCE
    sspec = importlib.util.spec_from_file_location(
        "sens", ROOT / "validators" / "run_spillover_sensitivity.py")
    sens = importlib.util.module_from_spec(sspec)
    sspec.loader.exec_module(sens)
    factor = EVIDENCE["spillover_share_pct_survey"].get("factor", 4.0)
    Z4, _ = sens.moved(Z, regions,
                       {j: ("ratio", factor) for j in range(len(regions))})
    E = np.zeros((len(X), S))
    E[np.arange(k * S, (k + 1) * S), np.arange(S)] = 1.0
    L4 = np.linalg.solve(np.eye(len(X)) - Z4 / np.where(X > 0, X, np.inf), E)
    m4 = L4.sum(0)
    intra4 = L4[s].sum(0)
    agg4 = float((m4 - intra4).sum() / m4.sum())
    per4 = (m4 - intra4) / m4
    check("and what the region would lose at the surveys' level of trade, "
          "from the counterfactual behind the archive-wide range",
          abs(ir.get("share_if_surveyed", -1.0) - agg4) < 1e-9
          and np.allclose(ir.get("share_by_sector_if_surveyed", []), per4,
                          atol=1e-9),
          f"{REGION}: {100 * agg:.1f} % as the archive has it, "
          f"{100 * agg4:.1f} % with its trade with the rest of the country "
          f"x{factor:g} — a counterfactual, not a corrected figure")

    check("every cell is marked ESTIMATED",
          t.provenance_counts() == {"ESTIMATED": S * S} and t.derived,
          "the archive is an estimate from the OECD ICIO, regional accounts "
          "and truck freight, not a survey, and a later split must not inherit "
          "it as a measurement")
    check("NPISH is not counted twice, and the notes say why",
          not any("NPISH" in l for l in t.Y_labels) and "NPISH" in (t.notes or ""),
          "NPISH and GGFC are the same column on all 2,720 rows")
    check("and the unit is the one the paper states, with its oddity",
          "dollar" in t.unit.lower(),
          f"{t.unit!r}")

    # ---- every refusal, fired
    from quadrium.io_loader import LoaderError  # noqa: F401  (named in msgs)

    check("the block's own Greek labels are not the ones the loader gives, so "
          "the test below means something",
          "EL11" in regions and "EL51" not in regions,
          "`load_Z` here reads the block's own labels, which print EL11 over "
          "Attiki's rows; the loader corrects them (run_mrio_labels.py)")
    for code, word, why in (
            ("UKI1", "no output", "empty in every file"),
            ("UKM3", "no output", "also empty; the count is four, not three"),
            ("FR10", "no other region", "Île-de-France trades with nobody in "
                                        "the archive, which no region does"),
            ("EL11", "EL51", "a NUTS 2010 code the archive prints over "
                             "another region's rows; the refusal says which "
                             "code to ask for instead (run_mrio_labels.py)"),
            ("XX99", "not in the archive", "a code nobody has")):
        ok, msg = refused(lambda c=code: load_eu_mrio_2018(MRIO, c), word)
        check(f"refused: {code} ({why})", ok, msg)

    # ---- the catalogue finds it, which is how a user learns it exists
    #
    # `--sources` promises every table the engine can load on this machine.
    # Until this was checked it did not look in data/mrio/ at all, so 259
    # regional tables were loadable and invisible.
    from quadrium.catalogue import advise, scan

    sources = scan(ROOT)
    # The 2018 tables: the catalogue lists every year it finds, and the
    # counts below are about one year's regions.
    mrio = [s for s in sources if s.table_kind == "eu_mrio" and s.year == 2018]
    empty = {r for j, r in enumerate(regions)
             if X[j * S:(j + 1) * S].sum() <= 0}
    mine = next((s for s in mrio if s.geo == REGION), None)
    check("the catalogue lists one regional table per region with output",
          len(mrio) == len(regions) - len(empty) and mine is not None
          and not empty & {s.geo for s in mrio}
          and mine.codes == list(t.sector_codes),
          f"{len(mrio)} of {len(regions)} regions; the {len(empty)} empty ones "
          f"are left out, read from the final-demand file and the header row "
          f"rather than the 35 MB block")
    a = advise("I55", sources, REGION)
    lines = a["best"]["source"].config_lines() if a.get("best") else []
    check("and `--find I55 --geo ES51` sends the user to divide `G-I` there",
          a["action"] == "split" and a["best"]["container"] == "G-I"
          and any(l.split()[:2] == ["mrio_region", REGION] for l in lines),
          " / ".join(l.strip() for l in lines) or a["why"][:160])
    nat = advise("I55", sources, "ES")
    check("while a question about Spain is still answered by a national table",
          nat.get("best") is not None
          and nat["best"]["source"].table_kind != "eu_mrio",
          nat["best"]["source"].source_id if nat.get("best")
          else nat["why"][:160])
    lonely = sorted({s.geo[:2] for s in mrio}
                    - {s.geo for s in sources if s.table_kind != "eu_mrio"})
    if lonely:
        g = lonely[0]
        why = advise("I55", sources, g)["why"]
        check(f"and a country with no national table here is pointed at its "
              f"regions ({g})",
              any(s.geo in why for s in mrio if s.geo[:2] == g)
              and "not the country" in why, why[:200])
    code, out, _ = run(["--sources", "--data", str(ROOT)])
    check("`--sources` says so in one paragraph rather than 268 lines",
          code == 0 and "European MRIO" in out
          and "mrio:eu2018" not in out, f"exit {code}")
    code, out, _ = run(["--find", "I55", "--geo", REGION, "--data", str(ROOT),
                        "--offline"])
    check("and `--find` prints the rows to paste, region included",
          code == 0 and "mrio_region" in out and REGION in out,
          f"exit {code}")

    # ---- from a workbook, which is how the owner will use it
    from quadrium.config import ConfigError, load_config, plan_workbook

    tmp = Path(tempfile.mkdtemp(prefix="quadrium_eumrio_"))
    base = [("project_id", "eumrio"), ("table_path", str(MRIO)),
            ("table_kind", "eu_mrio"), ("mrio_region", REGION)]
    good = workbook(tmp / "good.xlsx", base)
    code, out, err = run([str(good), "--outputs", str(tmp / "o")])
    rep = tmp / "o" / "eumrio" / "report.md"
    text = rep.read_text() if rep.exists() else ""
    check("a workbook naming `eu_mrio` and a region splits a sector of it",
          code == 0 and bool(text), f"exit {code}. {err.strip()[:160]}")
    check("and the report says the residue and what the table omits",
          "RESIDUAL" in text and f"{100 * agg:.1f} %" in text,
          "above every figure, where a reader meets it before the numbers")
    check("and shows it by sector, says where it goes, and that it is the "
          "table before the split",
          "What a one-region table leaves out" in text and top in text
          and "before any split" in text,
          "the run above divided G-I, and the figures are the archive's for "
          "the table as loaded: a split changes the region's own block, and "
          "the 2,720-sector inverse is not recomputed for the subsectors")
    check("and gives the surveys' level beside the archive's, labelled a "
          "counterfactual",
          "at the surveys' level" in text and f"{100 * agg4:.1f} %" in text
          and "counterfactual" in text,
          f"{100 * agg4:.1f} % for {REGION}, next to {100 * agg:.1f} %")
    moves = EVIDENCE["spillover_by_year"]["region_range_median_pts"]
    check("and says the region's figure is that year's, and how far such "
          "figures move across the deposit",
          "Read it as that year's" in text and f"{moves} points" in text,
          f"the archive's median barely moves across 2008-2018; a region's "
          f"own figure moves a median {moves} points "
          f"(run_spillover_years.py)")
    csv_path = tmp / "o" / "eumrio" / "scenarios" / "S1" / "table_disaggregated.csv"
    check("the residual column survives the split",
          csv_path.exists() and "RESIDUAL" in csv_path.read_text(),
          "a split divides rows and keeps the columns, so the column travels "
          "with the table rather than being dropped at the first operation")

    hfce = next((l for l in t.Y_labels if "HFCE" in l), "HFCE")
    vrow = next((l for l in t.VA_labels if "(VA)" in l), "VA")
    cases = (
        ("no region named", base[:3], "mrio_region"),
        ("a type II closure", base + [("type_ii_income_rows", vrow),
                                      ("type_ii_household_column", hfce)],
         "wages"),
        ("`mrio_region` on another kind",
         [("project_id", "x"), ("table_path", str(MRIO / FILES[0])),
          ("table_kind", "uk_analytical"), ("mrio_region", REGION)],
         "eu_mrio"),
    )
    for i, (what, meta, word) in enumerate(cases):
        wb = workbook(tmp / f"bad{i}.xlsx", meta)
        ok, msg = refused(lambda w=wb: load_config(w), word)
        check(f"the workbook is refused: {what}", ok, msg)

    reg = workbook(tmp / "reg.xlsx", base, splits=False,
                   regionalise=[("method", "FLQ"),
                                ("activity_path", str(tmp / "none.csv"))])
    ok, msg = refused(lambda: load_config(reg), "already")
    check("the workbook is refused: regionalising a table that is already "
          "regional", ok, msg)

    plan = plan_workbook(tmp / "bad0.xlsx")
    gaps = [g for g in plan["gaps"] if "mrio_region" in g["what"]]
    check("and --plan names the missing region before anything runs",
          not plan["ready"] and gaps and gaps[0]["severity"] == "blocking",
          gaps[0]["what"] if gaps else "no gap names it")
    plan = plan_workbook(tmp / "bad1.xlsx")
    check("and the type II request with it",
          any("type II" in g["what"] and g["severity"] == "blocking"
              for g in plan["gaps"]),
          "the archive publishes value added as one block with no wages row")

    print("\n" + "=" * 78)
    if FAIL:
        print(f"{len(FAIL)} check(s) FAILED:")
        for f in FAIL:
            print(f"  - {f}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
