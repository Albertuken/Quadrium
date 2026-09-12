"""
A region of the European MRIO, with its employment from Eurostat.

WHY
-----
The owner's first example for satellite accounts was an employment multiplier
for a region, and the engine could not give one: the accounts existed, but
somebody had to type the employment in. For a region of the European MRIO it
does not have to be typed. The archive's ten sectors are exactly Eurostat's A10
grouping, and Eurostat's regional accounts publish employment by NUTS-2 region
on that same grouping -- `nama_10r_3empers`, thousand persons, employed persons
`EMP`. So `mrio_employment: yes` fetches it for the region and year loaded,
keeps the download with its provenance, and attaches it as the account
`employment`.

WHAT IS CHECKED
-----------------
Catalonia, 2018, from the download kept in `data/eurostat/`, offline: the
account carries Eurostat's ten figures in the table's sector order, they add
up to the total Eurostat publishes, the report prints employment multipliers,
and the account's notes say what the ratio mixes -- employment is measured,
the archive's output is estimated.

WHAT IT CANNOT DO, AND SAYS
-----------------------------
Eurostat publishes the current NUTS codes and the archive codes France on NUTS
2013 and Greece on NUTS 2010, so there the codes do not meet. The engine
refuses and points at the `satellites` sheet rather than translating a code by
hand; `tests/test_engine.py` fires that refusal without a network.

AND THE SAME ACCOUNT ON THREE BLOCKS (2026-09-12)
---------------------------------------------------
`mrio_scope: with_rest` was refused with `mrio_employment` until then, because
an aggregate block's account had nothing to say about how much of the block
Eurostat covers. It says it now: each aggregate is the sum of Eurostat's own
figures for the regions inside it that Eurostat publishes -- 230 of the
archive's 268 -- and the share of the block's output those regions produce
travels with the account and is printed beside the figures. A block Eurostat
covers none of is refused rather than returned as zero.

What the account buys is a second landing table: where the JOBS an impulse
sets off land, which is not where the output lands. For Catalonia the widest
gap is real estate, 87.0 % of the jobs staying in the region against 93.9 %
of the output.

Run:
    python3 validators/run_eu_mrio_employment.py
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
import warnings
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
MRIO = ROOT / "data" / "mrio"
CUBE = ROOT / "data" / "eurostat" / "nama_10r_3empers_ALL_2018.json"
REGION, YEAR = "ES51", 2018
FAIL: list[str] = []

# Exit code 3, not 0. `check.sh` counts it as SKIPPED rather than as a pass:
# this validator opened no file and measured nothing, and a run that says
# "All checks passed" on evidence it never saw is the fault the suite exists
# to catch. It is not a failure -- a tree without the workbook still exits 0.
NOTHING_CHECKED = 3


def check(name: str, ok: bool, detail: str = "") -> None:
    print(f"  {'ok  ' if ok else 'FAIL'} {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        FAIL.append(name)


def cube_values(path, geo):
    """Eurostat's figure per NACE code for one region of the kept file, which
    holds every region the release carries."""
    from quadrium.eurostat import _Cube

    cube = _Cube(json.loads(Path(path).read_text()))
    return {code: cube.at(nace_r2=code, geo=geo, time=str(YEAR))
            for code in cube.index["nace_r2"]}


def main() -> int:
    warnings.filterwarnings("ignore")
    print(__doc__.strip().split("Run:")[0].rstrip())
    print("\n" + "=" * 78)

    need = [MRIO / f for f in ("MRIO_2018_272regions.xlsx",
                               "Final_demand_2018.xlsx", "TAXSUB_VA_2018.xlsx")]
    if not all(p.exists() for p in need) or not CUBE.exists():
        print("    -- the 2018 archive or the kept Eurostat download is absent.")
        print("\n" + "=" * 78)
        print("Nothing was checked: the archive or the employment download is "
              "not in this tree.")
        return NOTHING_CHECKED

    import openpyxl

    from quadrium.cli import main as cli
    from quadrium.config import load_config

    tmp = Path(tempfile.mkdtemp(prefix="quadrium_emp_"))
    cache = tmp / "data" / "eurostat"
    cache.mkdir(parents=True)
    shutil.copy(CUBE, cache / CUBE.name)
    side = CUBE.with_suffix(CUBE.suffix + ".provenance")
    if side.exists():
        shutil.copy(side, cache / side.name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "project"
    for k, v in (("project_id", "emp"), ("table_path", str(MRIO)),
                 ("table_kind", "eu_mrio"), ("mrio_region", REGION),
                 ("mrio_year", YEAR), ("mrio_employment", "sí")):
        ws.append([k, v])
    ws = wb.create_sheet("splits")
    ws.append(["sector_code", "new_code", "new_label", "key_id"])
    ws.append(["G-I", "GI1", "Trade and transport", "k1"])
    ws.append(["G-I", "GI2", "Accommodation and food", "k1"])
    ws = wb.create_sheet("keys")
    ws.append(["key_id", "new_sector_code", "value", "source", "source_year",
               "strength"])
    ws.append(["k1", "GI1", 70, "validator fixture, not a measurement", 2018,
               "weak"])
    ws.append(["k1", "GI2", 30, "validator fixture, not a measurement", 2018,
               "weak"])
    book = tmp / "emp.xlsx"
    wb.save(book)

    try:
        cfg = load_config(book, offline=True)
    except Exception as exc:                              # noqa: BLE001
        check("a workbook asking for the region's employment loads offline "
              "from the kept download", False, str(exc).splitlines()[0][:150])
        print(f"\n{len(FAIL)} check(s) FAILED.")
        return 1
    table = cfg["table"]
    sat = (table.satellites or {}).get("employment")
    want = cube_values(CUBE, REGION)
    expected = [want.get(c) for c in table.sector_codes]
    check("the account carries Eurostat's ten figures, in the table's sector "
          "order",
          sat is not None and None not in expected
          and list(sat.values) == expected,
          f"{REGION} {YEAR}: " + ", ".join(
              f"{c} {v:,.1f}" for c, v in zip(table.sector_codes, expected)
              if v is not None))
    check("and they add up to the total Eurostat publishes",
          sat is not None and abs(sum(sat.values) - want["TOTAL"]) < 0.5,
          f"{sum(sat.values) if sat else 0:,.1f} against a published "
          f"{want['TOTAL']:,.1f} thousand persons")
    check("and the account says what the ratio mixes",
          sat is not None and "measured" in (sat.notes or "")
          and "estimated" in (sat.notes or ""),
          (sat.notes or "")[:120] if sat else "no account")

    out = tmp / "out"
    with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
        code = cli([str(book), "--outputs", str(out), "--offline"])
    rep = out / "emp" / "report.md"
    text = rep.read_text() if rep.exists() else ""
    check("and the report prints employment multipliers for the region",
          code == 0 and "employment" in text and "thousand persons" in text,
          f"exit {code}")
    check("and says how much of them runs through other regions, sector by "
          "sector", "jobs through other regions" in text,
          "the employment version of what a one-region table leaves out")

    # ---- THE SAME ACCOUNT ON THREE BLOCKS. Refused until 2026-09-12 because
    # an aggregate's account had nothing to say about how much of the block
    # Eurostat covers. It says it now: each aggregate is the sum of the
    # figures for the regions inside it that Eurostat publishes, and the
    # coverage travels with it.
    wide_book = tmp / "emp_wide.xlsx"
    wb = openpyxl.load_workbook(book)
    wb["project"].append(["mrio_scope", "with_rest"])
    wb.save(wide_book)
    with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
        wcfg = load_config(wide_book, offline=True)
    wt = wcfg["table"]
    wsat = (wt.satellites or {}).get("employment")
    S = wt.n // 3
    cov = ((wt.interregional.get("jobs") or {}).get("coverage") or [])
    check("the account covers all thirty units of a three-block table",
          wsat is not None and len(wsat.values) == 3 * S
          and list(wsat.values[:S]) == expected[:S],
          f"{0 if wsat is None else len(wsat.values)} figures; the region's "
          f"own ten are the same ten as above")
    check("and each aggregate is Eurostat's own figures added up, with the "
          "share of the block it covers said",
          len(cov) == 2 and all(c["covered"] <= c["regions"] for c in cov)
          and any(c["covered"] < c["regions"] for c in cov)
          and "stated rather than filled" in (wsat.notes or ""),
          "; ".join(f"{c['block']}: {c['covered']} of {c['regions']} regions, "
                    f"{100 * c['output_share']:.1f} % of its output"
                    for c in cov))

    lands = wt.interregional["lands"]
    jl = (wt.interregional.get("jobs") or {}).get("lands") or []
    if jl:
        here_out = [row[0] / sum(row) for row in lands]
        here_job = [row[0] / sum(row) for row in jl]
        gap = [100 * (j - o) for j, o in zip(here_job, here_out)]
        worst = max(range(len(gap)), key=lambda i: abs(gap[i]))
        check("and a job and a euro of output do not land alike, which is why "
              "the account is worth attaching",
              max(abs(g) for g in gap) > 1.0,
              f"{wt.sector_codes[worst]}: {100 * here_job[worst]:.1f} % of "
              f"the jobs stay in {REGION} against "
              f"{100 * here_out[worst]:.1f} % of the output, "
              f"{gap[worst]:+.1f} points")

        guide = (ROOT / "docs" / "GUIDE.md").read_text()
        check("and the guide quotes that pair rather than a number that "
              "drifted",
              f"{100 * here_job[worst]:.1f} %" in guide
              and f"{100 * here_out[worst]:.1f} %" in guide,
              f"{100 * here_job[worst]:.1f} % and "
              f"{100 * here_out[worst]:.1f} % in docs/GUIDE.md")

    out2 = tmp / "out_wide"
    with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
        code2 = cli([str(wide_book), "--outputs", str(out2), "--offline"])
    rep2 = out2 / "emp" / "report.md"
    text2 = rep2.read_text() if rep2.exists() else ""
    check("and the report prints where the jobs land, block by block",
          code2 == 0 and "where the JOBS land" in text2
          and f"jobs in `{REGION}`" in text2,
          f"exit {code2}")

    print("\n" + "=" * 78)
    if FAIL:
        print(f"{len(FAIL)} check(s) FAILED:")
        for f in FAIL:
            print(f"  - {f}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
