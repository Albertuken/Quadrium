"""
The archive is in dollars, and every country of it is to scale.

WHY
-----
Two questions the project carried unanswered. The paper says the archive is in
"million dollars" in a sentence that also mentions environmental data the
deposit does not hold, so the loader's unit note has said "as Huang &
Koutroumpis (2023) state it" and left it there. And `run_mrio_labels.py` found
nineteen region labels naming another region by comparing a region's size with
Eurostat: the same comparison, taken to the country level, says whether any
country of the archive is out of scale.

Both fall out of one measurement. Eurostat publishes GDP by region
(`nama_10r_2gdp`), value added by region and branch (`nama_10r_3gva`) and the
year's average euro-dollar rate (`ert_bil_eur_a`), all kept in
`data/eurostat/` with their provenance.

WHAT IT FOUND
---------------
**The archive is in dollars.** The archive's value added over Eurostat's GDP is
1.313 in 2008 and 0.986 in 2015 -- no currency behaves like that. Divided by
the year's euro-dollar rate it is 0.890 in 2008 and 0.889 in 2015, and 0.885
to 0.892 in every one of the eleven years. A constant appears only in the right
currency.

**And that constant is what value added over GDP should be**: GDP is value
added plus taxes less subsidies on products, about a ninth. Over the 297
country-years the ratio is a median 0.890 (p10 0.842, p90 0.913).

**Each country's value added IS Eurostat's, converted.** For 2018, against
Eurostat's own value added rather than GDP, the archive gives a median 0.993
of it (p10 0.965, p90 1.014), from 0.893 for Malta to 1.038 for Latvia. So no
country is out of scale, and the archive's country totals are not an estimate
of their own.

**Inside the archive, a handful of units are out of line with themselves.**
Output more than ten times a unit's own value added means intermediate
consumption is nine tenths of production. Ten to fourteen of the 2,720 units
are like that in each year, nine of them in every year: agriculture in
metropolitan regions, where the archive leaves it almost no value added --
Berlin at 99 times, Brussels at 47, Madrid at 25, Vienna at 23. One more
joins them in 2013 and 2014, the construction of EL61 (Thessalia), and it is
the one jump the employment figures still have from year to year
(`run_employment_years.py`).

WHAT IT WOULD CATCH
---------------------
A country whose figures had been built on another scale -- or read in another
currency -- would sit apart from the rest in any year. None does; the whole
spread across 27 countries within a year is about six per cent.

Run:
    python3 validators/run_mrio_scale.py
"""
from __future__ import annotations

import json
import sys
import warnings
from collections import defaultdict
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
MRIO = ROOT / "data" / "mrio"
EUROSTAT = ROOT / "data" / "eurostat"
S = 10
FAIL: list[str] = []

# Exit code 3, not 0. `check.sh` counts it as SKIPPED rather than as a pass:
# this validator opened no file and measured nothing, and a run that says
# "All checks passed" on evidence it never saw is the fault the suite exists
# to catch. It is not a failure -- a tree without the workbook still exits 0.
NOTHING_CHECKED = 3


def check(name: str, ok: bool, detail: str = "") -> None:
    print(f"  {'ok  ' if ok else 'FAIL'} {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        FAIL.append(name)


def country_value_added(year: int):
    """The archive's value added by country, and how many of each country's
    regions it counts, from the side file and the corrected labels."""
    import openpyxl

    from quadrium.config import MRIO_REDRAWN
    from quadrium.io_loader import _mrio_files, _mrio_relabel, _mrio_side

    blk, _, vaf = _mrio_files(MRIO, year)
    wb = openpyxl.load_workbook(blk, read_only=True, data_only=True)
    try:
        head = next(wb.worksheets[0].iter_rows(values_only=True, max_row=1))
    finally:
        wb.close()
    regions = list(dict.fromkeys(
        l.split("-", 1)[0]
        for l in _mrio_relabel([str(x) for x in head[1:] if x is not None])))
    vh, VA = _mrio_side(vaf, "columns")
    va = VA[:, vh.index("VA")]
    got, counted, total = defaultdict(float), defaultdict(int), defaultdict(int)
    for k, r in enumerate(regions):
        total[r[:2]] += 1
        if r in MRIO_REDRAWN or r.startswith("UK"):
            continue
        got[r[:2]] += float(va[k * S:(k + 1) * S].sum())
        counted[r[:2]] += 1
    return got, counted, total, regions


def main() -> int:
    warnings.filterwarnings("ignore")
    print(__doc__.strip().split("Run:")[0].rstrip())
    print("\n" + "=" * 78)

    from quadrium.config import MRIO_EUROSTAT_CODE, MRIO_REDRAWN
    from quadrium.eurostat import _Cube
    from quadrium.io_loader import _MRIO_YEARS, _mrio_files

    years = []
    for y in _MRIO_YEARS:
        try:
            _mrio_files(MRIO, y)
        except Exception:                                 # noqa: BLE001
            continue
        if ((EUROSTAT / f"nama_10r_2gdp_ALL_{y}.json").exists()
                and (EUROSTAT / f"ert_bil_eur_a_USD_{y}.json").exists()):
            years.append(y)
    if not years:
        print("    -- no year has the archive, Eurostat's regional GDP and "
              "the euro-dollar rate in this tree.")
        print("\n" + "=" * 78)
        print("Nothing was checked: the archive or the Eurostat files are "
              "absent.")
        return NOTHING_CHECKED

    per_year_raw, per_year_usd, all_usd = {}, {}, []
    for y in years:
        got, counted, _total, regions = country_value_added(y)
        gdp = _Cube(json.loads(
            (EUROSTAT / f"nama_10r_2gdp_ALL_{y}.json").read_text()))
        rate = _Cube(json.loads(
            (EUROSTAT / f"ert_bil_eur_a_USD_{y}.json").read_text())
        ).at(currency="USD", time=str(y))
        euro = defaultdict(float)
        for r in regions:
            if r in MRIO_REDRAWN or r.startswith("UK"):
                continue
            v = gdp.at(geo=MRIO_EUROSTAT_CODE.get(r, r), time=str(y))
            if v is not None:
                euro[r[:2]] += float(v)
        raw = {c: got[c] / euro[c] for c in got if euro[c] > 0}
        per_year_raw[y] = float(np.median(list(raw.values())))
        usd = {c: v / rate for c, v in raw.items()}
        per_year_usd[y] = usd
        all_usd += list(usd.values())

    A = np.array(all_usd)
    med = [np.median(list(v.values())) for v in per_year_usd.values()]
    check("the archive's value added over Eurostat's GDP is a different "
          "number every year, so it is not in euro",
          max(per_year_raw.values()) - min(per_year_raw.values()) > 0.2,
          f"{min(per_year_raw.values()):.3f} to "
          f"{max(per_year_raw.values()):.3f} across {len(years)} years")
    check("and the same number every year once the year's euro-dollar rate is "
          "taken out, so it is in dollars",
          max(med) - min(med) < 0.02,
          f"median {min(med):.3f} to {max(med):.3f} across {len(years)} years")
    check("which is value added over GDP, about eight ninths, in every "
          "country and year",
          0.85 < float(np.median(A)) < 0.93
          and float(np.percentile(A, 90)) - float(np.percentile(A, 10)) < 0.12,
          f"{len(A)} country-years: median {np.median(A):.3f}, p10 "
          f"{np.percentile(A, 10):.3f}, p90 {np.percentile(A, 90):.3f}, "
          f"{A.min():.3f} to {A.max():.3f}")

    # ---- against Eurostat's own value added, where the branches are here
    gva_path = EUROSTAT / "nama_10r_3gva_ALL_2018.json"
    if 2018 in years and gva_path.exists():
        gva = _Cube(json.loads(gva_path.read_text()))
        rate = _Cube(json.loads(
            (EUROSTAT / "ert_bil_eur_a_USD_2018.json").read_text())
        ).at(currency="USD", time="2018")
        got, counted, total, _ = country_value_added(2018)
        rows = []
        for c in sorted(got):
            if counted[c] != total[c]:
                continue                  # a country with regions left out
            e = gva.at(nace_r2="TOTAL", geo=c, time="2018")
            if e:
                rows.append((got[c] / (float(e) * rate), c))
        v = np.array([x for x, _ in rows])
        check("and each country's value added IS Eurostat's, converted at "
              "that rate",
              len(rows) >= 20 and abs(float(np.median(v)) - 1.0) < 0.02
              and float(np.percentile(v, 10)) > 0.95
              and float(np.percentile(v, 90)) < 1.05,
              f"{len(rows)} countries whose regions are all counted: median "
              f"{np.median(v):.4f}, p10 {np.percentile(v, 10):.4f}, p90 "
              f"{np.percentile(v, 90):.4f}, {min(rows)[0]:.3f} "
              f"({min(rows)[1]}) to {max(rows)[0]:.3f} ({max(rows)[1]})")
    else:
        print("    -- Eurostat's value added by branch for 2018 is absent, so "
              "the comparison against value added itself was not made")

    # ---- the archive against itself: output far above a unit's value added
    #
    # No outside source is needed for this one. An output ten times the value
    # added means intermediate consumption is nine tenths of production, which
    # no industry sustains for long.
    import openpyxl

    from quadrium.io_loader import _MRIO_SECTORS, _mrio_relabel, _mrio_side

    sectors = list(_MRIO_SECTORS)
    out_of_line = {}
    for y in years:
        blk, fdf, vaf = _mrio_files(MRIO, y)
        wb = openpyxl.load_workbook(blk, read_only=True, data_only=True)
        try:
            head = next(wb.worksheets[0].iter_rows(values_only=True,
                                                   max_row=1))
        finally:
            wb.close()
        regions = list(dict.fromkeys(
            l.split("-", 1)[0]
            for l in _mrio_relabel([str(x) for x in head[1:]
                                    if x is not None])))
        fh, FD = _mrio_side(fdf, "rows")
        vh, VA = _mrio_side(vaf, "columns")
        X = FD[:, fh.index("TOTAL")].reshape(len(regions), S)
        V = VA[:, vh.index("VA")].reshape(len(regions), S)
        with np.errstate(invalid="ignore", divide="ignore"):
            ratio = np.where(V > 0, X / V, np.nan)
        out_of_line[y] = {f"{regions[i]}-{sectors[j]}": float(ratio[i, j])
                          for i, j in np.argwhere(np.isfinite(ratio)
                                                  & (ratio > 10))}
    counts = [len(v) for v in out_of_line.values()]
    everywhere = set.intersection(*[set(v) for v in out_of_line.values()])
    check("a handful of units carry an output more than ten times their own "
          "value added, and they are mostly the same ones every year",
          max(counts) <= 20 and len(everywhere) >= 8,
          f"{min(counts)} to {max(counts)} units a year of "
          f"{len(sectors) * 272:,}; {len(everywhere)} in every year: "
          f"{', '.join(sorted(everywhere)[:6])} — agriculture in metropolitan "
          f"regions, where the archive gives it almost no value added")
    el61 = sorted(y for y, v in out_of_line.items() if "EL61-F" in v)
    check("and the archive's construction in Thessalia is one of them in 2013 "
          "and 2014, which is the one jump the employment figures still have",
          el61 == [2013, 2014],
          f"EL61-F out of line in {el61 or 'no year'}: its output is 1,063 in "
          f"2012, 8,820 in 2013 and 710 in 2015 with its value added "
          f"unchanged (`run_employment_years.py`)")

    print("\n" + "=" * 78)
    if FAIL:
        print(f"{len(FAIL)} check(s) FAILED:")
        for f in FAIL:
            print(f"  - {f}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
